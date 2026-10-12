"""Opt-in end-to-end diagnostics for the frozen Windows release."""
import json
from pathlib import Path
import sys
import threading

from openpyxl import Workbook
from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pypdf import PdfReader

from app_version import VERSION
from diploma_core import Settings, generate


def self_test(directory):
    root = Path(directory).resolve()
    root.mkdir(parents=True, exist_ok=True)
    report = {"version": VERSION, "frozen": bool(getattr(sys, "frozen", False)),
              "success": False, "modes": {}}
    try:
        import wx
        from diploma_generator import DiplomaGeneratorApp
        app = wx.App(False)
        frame = DiplomaGeneratorApp(settings=Settings())
        frame.Layout()
        if frame.mode.GetCount() != 2:
            raise AssertionError("GUI mode switch missing")
        frame.timer.Stop()
        frame.Destroy()
        app.Yield()
        report["gui_constructed"] = True
        # Synthetic data only: never load the user's configuration or workbooks.
        image = Image.new("RGBA", (1600, 900), (250, 252, 255, 255))
        draw = ImageDraw.Draw(image)
        for x in range(0, image.width, 13):
            draw.line((x, 0, x, image.height), fill=(90, 130, 180, 255))
        image.save(root / "background.png")
        wb = Workbook()
        wb.active.append(["FullN1", "FullN2", "TEAMNAME", "CITY"])
        wb.active.append(["Тестовый участник с длинной фамилией Александр Константинович",
                          "Второй участник", "Тестовая команда", "Тюмень"])
        wb.save(root / "participants.xlsx")
        wb.close()
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(root / "background.png"), 0, 0, prs.slide_width, prs.slide_height)
        for top, token, font, size in [(1, "{NAME}", "Times New Roman", 32),
                                       (3, "{CITY}", "Arial", 20)]:
            shape = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(7), Inches(.7))
            shape.text = token
            run = shape.text_frame.paragraphs[0].runs[0]
            run.font.name, run.font.size = font, Pt(size)
            run.font.bold = True
        prs.save(root / "template.pptx")
        for mode, expected in (("row", 1), ("team", 2)):
            logs = []
            settings = Settings(str(root / "participants.xlsx"), str(root / "template.pptx"),
                                str(root / mode), mode=mode, auto_fit=True)
            count = generate(settings, logs.append, lambda *args: None, threading.Event())
            files = list((root / mode).glob("*.pdf"))
            if count != expected or len(files) != expected:
                raise AssertionError(f"{mode}: incorrect PDF count")
            for path in files:
                page = PdfReader(path).pages[0]
                if "Тюмень" not in page.extract_text() or "{NAME}" in page.extract_text():
                    raise AssertionError(f"{mode}: text substitution failed")
                actual = next(pic.image.convert("RGBA") for pic in page.images if pic.image.size == image.size)
                if ImageChops.difference(actual, image).getbbox(alpha_only=False) is not None:
                    raise AssertionError(f"{mode}: background pixels changed")
                font_names = [str(font.get_object().get("/BaseFont", ""))
                              for font in page["/Resources"]["/Font"].values()]
                if not any("TimesNewRoman" in font for font in font_names):
                    raise AssertionError(f"{mode}: template font missing")
            report["modes"][mode] = {"pdf_count": count, "background_pixel_exact": True,
                                     "selectable_text": True, "template_fonts": True,
                                     "text_shrunk": any("Текст уменьшен" in line for line in logs)}
        report["success"] = True
    except Exception as exc:
        import traceback
        report["error"] = str(exc)
        report["traceback"] = traceback.format_exc()
    (root / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["success"] else 1
