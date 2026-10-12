"""Data preparation and generation, independent of the desktop interface."""
from __future__ import annotations

import ctypes
import json
import logging
import multiprocessing as mp
import os
from pathlib import Path
import re
import shutil
import tempfile
import time
from dataclasses import dataclass, field
from datetime import date, datetime
import traceback

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

LOG = logging.getLogger(__name__)
PLACEHOLDER = re.compile(r"\{([^{}]+)\}")
MODES = {"row": "Один диплом на строку", "team": "Команда — несколько участников"}
POLICIES = {"stop": "Остановить", "skip": "Пропустить диплом", "default": "Заполнить по умолчанию"}


class DiplomaError(Exception):
    """An error that can be presented directly to the user."""


class Cancelled(Exception):
    pass


@dataclass
class Settings:
    excel_path: str = ""
    pptx_path: str = ""
    output_dir: str = ""
    mode: str = "team"
    column_mapping: dict[str, str] = field(default_factory=dict)
    error_handling: str = "stop"
    default_values: dict[str, str] = field(default_factory=dict)
    enable_sorting: bool = False
    sort_column: str = ""
    auto_fit: bool = True


@dataclass
class Record:
    row: int
    values: dict[str, str]


@dataclass
class Plan:
    headers: list[str]
    placeholders: list[str]
    records: list[Record]
    warnings: list[str]
    skipped: int = 0


def app_data_dir():
    base = Path(os.environ.get("LOCALAPPDATA", str(Path.home())))
    path = base / "DiplomaGenerator"
    path.mkdir(parents=True, exist_ok=True)
    return path


def absolute_path(value):
    value = str(value).strip().strip('"')
    if not value:
        raise DiplomaError("Не указан путь к файлу или папке.")
    return Path(os.path.abspath(os.path.expandvars(os.path.expanduser(value))))


def input_path(value, suffix):
    path = absolute_path(value)
    if path.suffix.lower() != suffix:
        raise DiplomaError(f"Нужен файл {suffix}: {path}")
    if not path.is_file():
        raise DiplomaError(f"Файл не найден: {path}\nПроверьте путь и доступность диска.")
    return path


def format_value(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def format_field(placeholder, value):
    text = format_value(value)
    if placeholder.upper() in ("DATE", "ДАТА") and text:
        for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, pattern).strftime("%d.%m.%Y")
            except ValueError:
                continue
    return text


def read_excel(value, headers_only=False, stop_event=None):
    path = input_path(value, ".xlsx")
    wb = formula_wb = None
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        if ws is None or ws.max_row < 1:
            raise DiplomaError("В Excel нет активного листа с данными.")
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [format_value(x).strip("{}").strip() or f"Столбец_{get_column_letter(i)}"
                   for i, x in enumerate(first, 1)]
        if len(set(headers)) != len(headers):
            duplicates = sorted({x for x in headers if headers.count(x) > 1})
            raise DiplomaError("Повторяются заголовки Excel: " + ", ".join(duplicates))
        if headers_only:
            return headers, []
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if stop_event and stop_event.is_set():
                raise Cancelled()
            rows.append(row)
        formula_wb = load_workbook(path, read_only=True, data_only=False)
        formulas = formula_wb[ws.title].iter_rows(min_row=2)
        for row_index, (row, formula_row) in enumerate(zip(rows, formulas), 2):
            if stop_event and stop_event.is_set():
                raise Cancelled()
            for i, cell in enumerate(formula_row):
                if cell.data_type == "f" and row[i] is None:
                    raise DiplomaError(f"Ячейка {cell.coordinate}: у формулы нет сохранённого результата. "
                                       "Откройте Excel, пересчитайте книгу и сохраните её.")
                if cell.data_type == "e":
                    raise DiplomaError(f"Ячейка {cell.coordinate}: ошибка Excel {cell.value}.")
        return headers, rows
    except (DiplomaError, Cancelled):
        raise
    except PermissionError as exc:
        raise DiplomaError(f"Нет доступа к Excel: {path}. Закройте файл или проверьте права.") from exc
    except Exception as exc:
        raise DiplomaError(f"Не удалось прочитать Excel: {path}\n{exc}") from exc
    finally:
        if wb:
            wb.close()
        if formula_wb:
            formula_wb.close()


def frame_targets(shapes, path=()):
    for index, shape in enumerate(shapes, 1):
        current = path + (index,)
        if shape.has_text_frame:
            yield shape.text_frame, {"path": current, "cell": None, "name": shape.name}
        if shape.has_table:
            for row_index, row in enumerate(shape.table.rows, 1):
                for column_index, cell in enumerate(row.cells, 1):
                    if not cell.is_spanned:
                        yield cell.text_frame, {"path": current, "cell": (row_index, column_index),
                                                "name": f"{shape.name}, ячейка {row_index}:{column_index}"}
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from frame_targets(shape.shapes, current)


def text_frames(shapes):
    for frame, _ in frame_targets(shapes):
        yield frame


def read_template(value):
    path = input_path(value, ".pptx")
    try:
        prs = Presentation(str(path))
        if not prs.slides:
            raise DiplomaError("В шаблоне нет слайдов.")
        placeholders = sorted({m.group(1) for frame in text_frames(prs.slides[0].shapes)
                               for p in frame.paragraphs for m in PLACEHOLDER.finditer(p.text)})
        if not placeholders:
            raise DiplomaError("На первом слайде нет полей вида {NAME}, {DATE} или {ФИО}.")
        return prs, placeholders
    except DiplomaError:
        raise
    except Exception as exc:
        raise DiplomaError(f"Не удалось прочитать шаблон: {path}\n{exc}") from exc


SYNONYMS = {
    "NAME": ["фио", "имя", "name", "full name", "участник"],
    "ORG": ["организация", "учреждение", "школа", "organization"],
    "TEAM": ["teamname", "команда", "team", "название команды"],
    "TEAMNAME": ["команда", "team", "название команды"],
    "CAT": ["категория", "category", "номинация", "возраст"],
    "TRAINER": ["тренер", "руководитель", "teacher", "coach"],
    "DATE": ["дата", "date"], "PLACE": ["место", "place", "position", "rank"],
    "CITY": ["город", "city", "населенный пункт", "location"],
}


def auto_mapping(placeholders, headers, mode):
    result = {}
    for ph in placeholders:
        if mode == "team" and ph in ("NAME", "TRAINER"):
            continue
        exact = next((h for h in headers if h.casefold() == ph.casefold()), None)
        if exact:
            result[ph] = exact
            continue
        aliases = SYNONYMS.get(ph.upper(), [])
        match = next((h for h in headers if h.casefold() in aliases), None)
        if not match and ph.upper() == "NAME":
            match = next((h for h in headers if h.casefold().startswith("fulln")), None)
        if not match and ph.upper() == "TRAINER":
            match = next((h for h in headers if h.casefold().startswith("fulltr")), None)
        if match:
            result[ph] = match
    return result


def prepare_plan(settings, stop_event=None):
    if settings.mode not in MODES or settings.error_handling not in POLICIES:
        raise DiplomaError("Неизвестный режим генерации или обработки пустых полей.")
    headers, rows = read_excel(settings.excel_path, stop_event=stop_event)
    prs, placeholders = read_template(settings.pptx_path)
    warnings = []
    if len(prs.slides) > 1:
        warnings.append("В шаблоне несколько слайдов: используется только первый.")
    mapping = auto_mapping(placeholders, headers, settings.mode)
    mapping.update(settings.column_mapping)
    for ph, header in mapping.items():
        if header and header not in headers:
            raise DiplomaError(f"Поле {{{ph}}} связано с отсутствующим столбцом «{header}». Обновите сопоставление.")
    people = [i for i, h in enumerate(headers) if h.casefold().startswith("fulln")]
    trainers = [i for i, h in enumerate(headers) if h.casefold().startswith("fulltr")]
    if settings.mode == "team" and not people:
        raise DiplomaError("Для режима команд нужны столбцы FullN1, FullN2… "
                           "Для таблицы с одной колонкой ФИО выберите «Один диплом на строку».")
    records, skipped, issues = [], 0, []
    for row_index, row in enumerate(rows, 2):
        if stop_event and stop_event.is_set():
            raise Cancelled()
        if not any(format_value(x) for x in row):
            continue
        base = {h: format_value(row[i]) for i, h in enumerate(headers)}
        if settings.mode == "team":
            names = [format_value(row[i]) for i in people if format_value(row[i])]
            trainer = next((format_value(row[i]) for i in trainers if format_value(row[i])), "")
            if not names:
                skipped += 1
                warnings.append(f"Строка {row_index}: нет участников, строка пропущена.")
                continue
        else:
            names, trainer = [None], ""
        for name in names:
            values = dict(base)
            for ph in placeholders:
                header = mapping.get(ph, "")
                values[ph] = format_field(ph, base.get(header, "")) if header else ""
            if settings.mode == "team":
                values["NAME"] = name
                values["TRAINER"] = trainer
            missing = [ph for ph in placeholders if not values.get(ph)]
            if settings.error_handling == "default":
                for ph in missing:
                    values[ph] = format_field(ph, settings.default_values.get(ph))
                missing = [ph for ph in missing if not values.get(ph)]
            if missing:
                issue = f"Строка {row_index}" + (f", {name}" if name else "") + ": пустые поля " + ", ".join("{" + ph + "}" for ph in missing)
                if settings.error_handling == "skip":
                    skipped += 1
                    warnings.append(issue + "; диплом пропущен.")
                    continue
                issues.append(issue)
                continue
            if settings.mode == "row" and not values.get("NAME"):
                name_header = settings.column_mapping.get("NAME") or auto_mapping(["NAME"], headers, "row").get("NAME")
                values["NAME"] = base.get(name_header, "") or f"Диплом_{row_index}"
            records.append(Record(row_index, values))
    if issues:
        extra = f"\n…ещё ошибок: {len(issues)-8}" if len(issues) > 8 else ""
        raise DiplomaError("Заполните данные, настройте сопоставление или выберите другой режим пустых полей:\n" + "\n".join(issues[:8]) + extra)
    if not records:
        raise DiplomaError("Нет дипломов для генерации. Проверьте участников и правила пустых полей.\n" + "\n".join(warnings[:5]))
    if settings.enable_sorting:
        if not settings.sort_column:
            raise DiplomaError("Выберите поле для распределения по папкам.")
        if not any(settings.sort_column in record.values for record in records):
            raise DiplomaError(f"Поле для папок «{settings.sort_column}» отсутствует в данных.")
    return Plan(headers, placeholders, records, warnings, skipped)


def replace_paragraph(paragraph, values):
    """Replace all tokens once, preserving run styles, including split tokens."""
    full_text = paragraph.text
    matches = list(PLACEHOLDER.finditer(full_text))
    if not matches:
        return
    # Include soft line breaks in offsets, matching python-pptx paragraph.text.
    offsets, pos, run_index = [], 0, 0
    runs = list(paragraph.runs)
    for child in paragraph._p:
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "r":
            run = runs[run_index]
            run_index += 1
            offsets.append((run, pos, pos + len(run.text)))
            pos += len(run.text)
        elif tag == "br":
            pos += 1
        elif tag == "fld":
            pos += len("".join(child.itertext()))
    for match in reversed(matches):
        replacement = format_value(values.get(match.group(1), ""))
        touched = [(run, start, end) for run, start, end in offsets
                   if start < match.end() and end > match.start()]
        if not touched:
            continue
        for index, (run, start, end) in enumerate(touched):
            left = max(0, match.start() - start)
            right = min(end - start, match.end() - start)
            run.text = run.text[:left] + (replacement if index == 0 else "") + run.text[right:]


def fill_presentation(prs, values, auto_fit=True, font_settings=None):
    targets = []
    for frame, target in frame_targets(prs.slides[0].shapes):
        changed = False
        for paragraph in frame.paragraphs:
            if PLACEHOLDER.search(paragraph.text):
                replace_paragraph(paragraph, values)
                changed = True
                if font_settings and font_settings.get("use_custom"):
                    for run in paragraph.runs:
                        run.font.name = font_settings.get("name", "Arial")
                        run.font.size = Pt(font_settings.get("size", 20))
                        run.font.bold = font_settings.get("bold", False)
        if changed and auto_fit:
            frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            frame.word_wrap = False
            targets.append(target)
    # The original application exports only the first slide.
    for slide_id in list(prs.slides._sldIdLst)[1:]:
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    return targets


def fit_powerpoint_text(deck, targets):
    """Measure with Office itself and shrink font sizes without changing styles.

    PowerPoint does not reliably recalculate normAutofit after external XML edits.
    Explicit paragraph and soft line breaks remain; automatic wrapping is disabled.
    """
    results = []
    for target in targets:
        shapes = deck.Slides.Item(1).Shapes
        for index in target["path"][:-1]:
            shapes = shapes.Item(index).GroupItems
        shape = shapes.Item(target["path"][-1])
        if target["cell"]:
            row, col = target["cell"]
            shape = shape.Table.Cell(row, col).Shape
        frame = shape.TextFrame2
        # Measure at the real font size instead of a cached AutoFit scale.
        # PowerPoint rejects AutoSize on table cells; manual fitting still works.
        is_cell = bool(target["cell"])
        if not is_cell:
            frame.AutoSize = 0
            frame.WordWrap = 0
        text = frame.TextRange
        if not text.Text.strip():
            if not is_cell:
                frame.AutoSize = 2
            continue
        width = float(shape.Width) - float(frame.MarginLeft) - float(frame.MarginRight)
        height = float(shape.Height) - float(frame.MarginTop) - float(frame.MarginBottom)
        if width <= 0 or height <= 0:
            raise DiplomaError(f"Блок «{target['name']}»: отступы не оставляют места для текста.")
        runs = [text.Runs(i, 1) for i in range(1, text.Runs().Count + 1)]
        original_sizes = [float(run.Font.Size) for run in runs]
        if any(size <= 0 for size in original_sizes):
            raise DiplomaError(f"Блок «{target['name']}»: не удалось определить размер шрифта.")
        scale = 1.0
        expected_lines = len(re.split(r"\r\n|[\r\n\v]", text.Text))
        for attempt in range(24):
            bound_width, bound_height = float(text.BoundWidth), float(text.BoundHeight)
            actual_lines = int(text.Lines().Count)
            if (bound_width <= width + 0.2 and bound_height <= height + 0.2
                    and actual_lines <= expected_lines):
                break
            measured_width = bound_width
            if is_cell and actual_lines > expected_lines:
                # Table cells force WordWrap and reject the COM setter. Sum the
                # widths of their rendered fragments to measure each intended line.
                for match in re.finditer(r"[^\r\n\v]+", text.Text):
                    line = text.Characters(match.start() + 1, match.end() - match.start())
                    measured_width = max(measured_width,
                                         sum(float(line.Lines(i, 1).BoundWidth)
                                             for i in range(1, line.Lines().Count + 1)))
            factor = min(width / max(measured_width, 0.01), height / max(bound_height, 0.01), 0.98) * 0.98
            scale *= factor
            sizes = [max(1.0, round(size * scale, 2)) for size in original_sizes]
            for run, size in zip(runs, sizes):
                run.Font.Size = size
        else:
            raise DiplomaError(f"Блок «{target['name']}»: текст не помещается даже после уменьшения. "
                               "Увеличьте блок или уменьшите его отступы и интервалы в шаблоне.")
        # Keep the PowerPoint 'Shrink text on overflow' setting as well.
        if not is_cell:
            frame.AutoSize = 2
        results.append({"name": target["name"], "scale": scale,
                        "width": float(text.BoundWidth), "available_width": width,
                        "height": float(text.BoundHeight), "available_height": height,
                        "lines": int(text.Lines().Count)})
    return results


def safe_component(value, fallback="Без значения", limit=80):
    text = re.sub(r'[\x00-\x1f<>:"/\\|?*]', "_", format_value(value)).strip(" .")
    text = text[:limit].rstrip(" .") or fallback
    if re.fullmatch(r"(?i)(CON|PRN|AUX|NUL|COM[1-9¹²³]|LPT[1-9¹²³])(?:\..*)?", text):
        text = "_" + text
    return text


def validate_output(value):
    path = absolute_path(value)
    try:
        path.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryFile(dir=path):
            pass
    except OSError as exc:
        raise DiplomaError(f"Нельзя записать файлы в папку: {path}\nПроверьте права, свободное место и доступность диска.\n{exc}") from exc
    return path


def publish_pdf(source, directory, stem):
    """Reserve a unique destination, never overwrite an existing diploma."""
    directory.mkdir(parents=True, exist_ok=True)
    for index in range(1, 100000):
        suffix = "" if index == 1 else f" ({index})"
        destination = directory / f"{stem}{suffix}.pdf"
        # Office still has path-length restrictions even on long-path Windows.
        if len(str(destination)) > 240:
            raise DiplomaError(f"Слишком длинный путь к PDF ({len(str(destination))} символов). "
                               f"Выберите папку ближе к корню диска: {directory}")
        try:
            target = destination.open("xb")
        except FileExistsError:
            continue
        try:
            with target, Path(source).open("rb") as src:
                shutil.copyfileobj(src, target)
        except BaseException:
            destination.unlink(missing_ok=True)
            raise
        return destination
    raise DiplomaError("Слишком много дипломов с одинаковым именем.")


def _powerpoint_worker(input_pptx, output_pdf, connection, fit_targets):
    """Isolated COM apartment. Report exactly the Office process we own."""
    import comtypes
    import comtypes.client
    import psutil
    app = deck = None
    owned = False
    initialized = False
    stage = "инициализация Office"
    try:
        comtypes.CoInitialize()
        initialized = True
        existing = {p.pid for p in psutil.process_iter(["name"])
                    if (p.info["name"] or "").lower() == "powerpnt.exe"}
        # HWND is a hidden typelib property and is unavailable via dynamic IDispatch.
        app = comtypes.client.CreateObject("PowerPoint.Application")
        stage = "определение процесса PowerPoint"
        pid = ctypes.c_ulong()
        ctypes.windll.user32.GetWindowThreadProcessId(ctypes.c_void_p(int(app.HWND)), ctypes.byref(pid))
        owned = bool(pid.value and pid.value not in existing)
        if owned:
            connection.send(("owned", (pid.value, psutil.Process(pid.value).create_time())))
        stage = "открытие временного шаблона"
        deck = app.Presentations.Open(str(input_pptx), True, False, False)
        if fit_targets:
            stage = "уменьшение текста до размеров блоков"
            report = fit_powerpoint_text(deck, fit_targets)
            connection.send(("fitted", report))
        stage = "сохранение PDF"
        # Explicit print intent. PNG backgrounds can still be recompressed by
        # Office, so their original pixels are restored after the worker closes.
        deck.ExportAsFixedFormat(str(output_pdf), 2, 2, PrintRange=None)
        deck.Close()
        deck = None
        connection.send(("success", None))
    except BaseException as exc:
        LOG.exception("PowerPoint failed: %s", stage)
        try:
            connection.send(("error", {"message": f"Этап: {stage}. {exc}",
                                       "traceback": traceback.format_exc()}))
        except (BrokenPipeError, OSError):
            pass
    finally:
        if deck is not None:
            try:
                deck.Close()
            except Exception:
                LOG.exception("Could not close generated presentation")
        if app is not None and owned:
            try:
                app.Quit()
            except Exception:
                LOG.exception("Could not quit owned PowerPoint")
        app = deck = None
        if initialized:
            comtypes.CoUninitialize()
        connection.close()


def _terminate_owned(identity):
    if not identity:
        return
    import psutil
    pid, creation_time = identity
    try:
        process = psutil.Process(pid)
        if process.create_time() == creation_time and process.name().lower() == "powerpnt.exe":
            process.terminate()
            try:
                process.wait(3)
            except psutil.TimeoutExpired:
                process.kill()
    except (psutil.NoSuchProcess, psutil.AccessDenied):
        LOG.exception("Could not terminate owned PowerPoint")


def pptx_to_pdf(input_pptx, output_pdf, stop_event, timeout=90, fit_targets=None, fit_report=None,
                quality_report=None):
    if stop_event.is_set():
        raise Cancelled()
    if os.name != "nt":
        raise DiplomaError("Экспорт PDF поддерживается в Windows с установленным Microsoft PowerPoint.")
    context = mp.get_context("spawn")
    reader, writer = context.Pipe(duplex=False)
    process = context.Process(target=_powerpoint_worker,
                              args=(str(Path(input_pptx).resolve()), str(Path(output_pdf).resolve()), writer, fit_targets or []))
    identity, outcome = None, None
    start = time.monotonic()
    started = False
    try:
        process.start()
        started = True
        writer.close()
        while True:
            if stop_event.is_set():
                raise Cancelled()
            if time.monotonic() - start > timeout:
                raise DiplomaError(f"PowerPoint не завершил экспорт за {timeout} секунд. "
                                   "Проверьте, нет ли диалога активации Office или ошибки открытия шаблона.")
            if reader.poll(0.1):
                try:
                    kind, payload = reader.recv()
                except EOFError:
                    break
                if kind == "owned":
                    identity = payload
                elif kind == "fitted":
                    LOG.info("Text fitting: %s", payload)
                    if fit_report:
                        fit_report(payload)
                else:
                    outcome = (kind, payload)
                    break
            if not process.is_alive():
                if reader.poll():
                    continue
                break
        if not outcome or outcome[0] != "success":
            detail = outcome[1] if outcome else f"Процесс экспорта завершился: {process.exitcode}"
            if isinstance(detail, dict):
                LOG.error("Office traceback:\n%s", detail.get("traceback", ""))
                detail = detail["message"]
            raise DiplomaError("Не удалось экспортировать PDF через PowerPoint. "
                               "Убедитесь, что Microsoft PowerPoint установлен и активирован.\n" + str(detail))
        pdf = Path(output_pdf)
        if not pdf.is_file() or pdf.stat().st_size < 5:
            raise DiplomaError("PowerPoint не создал корректный PDF.")
        with pdf.open("rb") as stream:
            if stream.read(5) != b"%PDF-":
                raise DiplomaError("PowerPoint не создал корректный PDF.")
        if stop_event.is_set():
            raise Cancelled()
        from pdf_quality import restore_original_images
        quality = restore_original_images(input_pptx, output_pdf)
        LOG.info("PDF image quality: %s", quality)
        if quality_report:
            quality_report(quality)
        return True
    finally:
        if started:
            process.join(timeout=2)
            if process.is_alive():
                process.terminate()
                process.join(timeout=3)
            _terminate_owned(identity)
            process.close()
        reader.close()
        writer.close()


def generate(settings, log, progress, stop_event, converter=None, plan=None):
    converter = converter or pptx_to_pdf
    plan = plan or prepare_plan(settings, stop_event)
    output = validate_output(settings.output_dir)
    for warning in plan.warnings[:30]:
        log(warning)
    log(f"Режим: {MODES[settings.mode]}. Дипломов: {len(plan.records)}. Пропущено: {plan.skipped}.")
    # A private short temporary directory prevents clashes and Office path issues.
    started, completed = time.monotonic(), 0
    with tempfile.TemporaryDirectory(prefix="diploma_") as temp:
        temp_dir = Path(temp)
        for index, record in enumerate(plan.records, 1):
            if stop_event.is_set():
                raise Cancelled()
            try:
                prs, _ = read_template(settings.pptx_path)
                fit_targets = fill_presentation(prs, record.values, settings.auto_fit)
                pptx, pdf = temp_dir / "document.pptx", temp_dir / "document.pdf"
                pdf.unlink(missing_ok=True)
                prs.save(str(pptx))
                log(f"{index}/{len(plan.records)} · {record.values.get('NAME', '')}: экспорт PDF…")
                def report_fit(report):
                    for item in report:
                        if item["scale"] < 0.999:
                            log(f"Текст уменьшен в блоке «{item['name']}»: {item['scale']:.0%} от исходного размера.")
                if converter is pptx_to_pdf:
                    def report_quality(report):
                        if report["restored"]:
                            log(f"Исходное качество картинок восстановлено без потерь: {len(report['restored'])}.")
                        if report["unmatched"]:
                            log("Часть картинок не удалось однозначно сопоставить с оригиналом; "
                                "для них сохранён экспорт Office: " + ", ".join(report["unmatched"]))
                    converted = converter(str(pptx), str(pdf), stop_event,
                                          fit_targets=fit_targets, fit_report=report_fit,
                                          quality_report=report_quality)
                else:
                    converted = converter(str(pptx), str(pdf), stop_event)
                if converted is False:
                    raise Cancelled()
                if stop_event.is_set():
                    raise Cancelled()
                directory = output
                if settings.enable_sorting:
                    directory /= safe_component(record.values.get(settings.sort_column, ""))
                name = safe_component(record.values.get("NAME"), f"Диплом_{record.row}")
                team = record.values.get("TEAMNAME") or record.values.get("TEAM", "")
                stem = safe_component(team, limit=50) + "_" + name if team and settings.mode == "team" else name
                destination = publish_pdf(pdf, directory, stem)
                completed += 1
                log(f"Сохранён: {destination}")
                elapsed = time.monotonic() - started
                eta = int(elapsed / completed * (len(plan.records) - completed))
                progress(completed, len(plan.records), eta)
            except Cancelled:
                raise
            except Exception as exc:
                LOG.exception("Generation failed at Excel row %s", record.row)
                raise DiplomaError(f"Строка {record.row}, {record.values.get('NAME', '')}: {exc}\n"
                                   f"Уже сохранено дипломов: {completed}. Папка: {output}") from exc
    return completed


def save_settings(settings, path=None):
    from dataclasses import asdict
    path = Path(path) if path else app_data_dir() / "config.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent,
                                     suffix=".tmp", delete=False) as stream:
        temporary = Path(stream.name)
        json.dump(asdict(settings), stream, ensure_ascii=False, indent=2)
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def load_settings(path=None):
    path = Path(path) if path else app_data_dir() / "config.json"
    if not path.exists():
        return Settings()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError("Ожидался объект JSON")
        settings = Settings()
        for key in ("excel_path", "pptx_path", "output_dir", "mode", "error_handling", "sort_column"):
            if key in data:
                if not isinstance(data[key], str):
                    raise ValueError(f"{key}: ожидалась строка")
                setattr(settings, key, data[key])
        for key in ("column_mapping", "default_values"):
            if key in data:
                if not isinstance(data[key], dict) or not all(isinstance(k, str) and isinstance(v, str) for k, v in data[key].items()):
                    raise ValueError(f"{key}: ожидался словарь строк")
                setattr(settings, key, data[key])
        for key in ("enable_sorting", "auto_fit"):
            if key in data:
                if not isinstance(data[key], bool):
                    raise ValueError(f"{key}: ожидалось логическое значение")
                setattr(settings, key, data[key])
        if settings.mode not in MODES or settings.error_handling not in POLICIES:
            raise ValueError("Неизвестный режим")
        return settings
    except (ValueError, OSError, TypeError) as exc:
        raise DiplomaError(f"Не удалось загрузить настройки {path}: {exc}") from exc
