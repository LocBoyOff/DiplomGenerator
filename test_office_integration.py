"""Opt-in integration checks against installed Microsoft PowerPoint."""
import os
from pathlib import Path
import tempfile
import threading
import unittest

from openpyxl import Workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from diploma_core import (Settings, fill_presentation, fit_powerpoint_text, generate,
                          prepare_plan)


@unittest.skipUnless(os.environ.get('RUN_OFFICE_TESTS') == '1', 'Set RUN_OFFICE_TESTS=1 to use PowerPoint')
class OfficeIntegrationTests(unittest.TestCase):
    def setUp(self):
        import comtypes.client
        import ctypes
        import psutil
        self.temp = tempfile.TemporaryDirectory(prefix='diploma_office_')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        before = {p.pid for p in psutil.process_iter(['name'])
                  if (p.info['name'] or '').lower() == 'powerpnt.exe'}
        self.app = comtypes.client.CreateObject('PowerPoint.Application')
        pid = ctypes.c_ulong()
        ctypes.windll.user32.GetWindowThreadProcessId(ctypes.c_void_p(int(self.app.HWND)), ctypes.byref(pid))
        self.owned = pid.value not in before
        self.addCleanup(self.close_app)

    def close_app(self):
        if self.owned and self.app.Presentations.Count == 0:
            self.app.Quit()
        self.app = None

    def fixture(self):
        name = 'Сверхдлинная-Фамилия-Участника Александр Константинович'
        wb = Workbook()
        wb.active.append(['FullN1', 'CITY'])
        wb.active.append([name, 'Тюмень'])
        excel = self.root / 'Таблица.xlsx'
        wb.save(excel)
        wb.close()
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for i, (token, font, size) in enumerate([('{NAME}', 'Times New Roman', 32), ('{CITY}', 'Georgia', 20)]):
            shape = slide.shapes.add_textbox(Inches(1), Inches(1 + i * 2), Inches(3), Inches(1.5))
            shape.name = 'ФИО' if i == 0 else 'Город'
            p = shape.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = token[:3]
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(30, 70, 150)
            run2 = p.add_run()
            run2.text = token[3:]
            run2.font.name = font
            run2.font.size = Pt(size)
        template = self.root / 'Шаблон.pptx'
        prs.save(template)
        return excel, template, name

    def test_fonts_shrink_and_export_in_both_modes(self):
        excel, template, name = self.fixture()
        for mode in ('row', 'team'):
            with self.subTest(mode=mode):
                settings = Settings(str(excel), str(template), str(self.root / mode), mode=mode)
                plan = prepare_plan(settings)
                prs = Presentation(str(template))
                targets = fill_presentation(prs, plan.records[0].values)
                path = self.root / (mode + '.pptx')
                prs.save(path)
                deck = self.app.Presentations.Open(str(path), False, False, False)
                try:
                    shape = deck.Slides.Item(1).Shapes.Item(1)
                    city = deck.Slides.Item(1).Shapes.Item(2)
                    font_before = shape.TextFrame2.TextRange.Runs(1, 1).Font
                    expected = (font_before.Name, font_before.Bold, font_before.Italic,
                                font_before.Fill.ForeColor.RGB)
                    report = fit_powerpoint_text(deck, targets)
                    self.assertEqual(shape.TextFrame2.TextRange.Text, name)
                    font_after = shape.TextFrame2.TextRange.Runs(1, 1).Font
                    self.assertEqual((font_after.Name, font_after.Bold, font_after.Italic,
                                      font_after.Fill.ForeColor.RGB), expected)
                    self.assertEqual(font_after.Name, 'Times New Roman')
                    self.assertLess(font_after.Size, 32)
                    self.assertEqual(city.TextFrame2.TextRange.Runs(1, 1).Font.Name, 'Georgia')
                    self.assertEqual(city.TextFrame2.TextRange.Runs(1, 1).Font.Size, 20)
                    self.assertEqual(report[0]['lines'], 1)
                    for item in report:
                        self.assertLessEqual(item['width'], item['available_width'] + .2)
                        self.assertLessEqual(item['height'], item['available_height'] + .2)
                    deck.SaveAs(str(self.root / (mode + '_fitted.pdf')), 32)
                    print(mode, 'FIT_REPORT', report)
                finally:
                    deck.Close()
                # Also exercise the isolated worker used by the desktop app.
                count = generate(settings, lambda x: None, lambda *x: None, threading.Event())
                self.assertEqual(count, 1)
                self.assertEqual(len(list((self.root / mode).glob('*.pdf'))), 1)

    def test_mixed_fonts_groups_table_and_explicit_line_breaks(self):
        excel, template, name = self.fixture()
        prs = Presentation(str(template))
        slide = prs.slides[0]
        group = slide.shapes.add_group_shape()
        box = group.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
        paragraph = box.text_frame.paragraphs[0]
        for value, font, size in [('Имя: ', 'Arial', 18), ('{NAME}', 'Times New Roman', 30),
                                  (' · участник', 'Georgia', 20)]:
            run = paragraph.add_run()
            run.text = value
            run.font.name, run.font.size = font, Pt(size)
        table = slide.shapes.add_table(1, 1, Inches(5), Inches(1), Inches(4), Inches(1)).table
        table.cell(0, 0).text = '{NAME}'
        table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.name = 'Arial'
        table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size = Pt(22)
        lines = slide.shapes.add_textbox(Inches(5), Inches(3), Inches(4), Inches(1.5))
        lines.text = 'Диплом\n{NAME}'
        for p in lines.text_frame.paragraphs:
            for run in p.runs:
                run.font.name, run.font.size = 'Georgia', Pt(24)
        targets = fill_presentation(prs, {'NAME': name, 'CITY': 'Тюмень'})
        path = self.root / 'complex.pptx'
        prs.save(path)
        deck = self.app.Presentations.Open(str(path), False, False, False)
        try:
            mixed = deck.Slides.Item(1).Shapes.Item(3).GroupItems.Item(1).TextFrame2.TextRange
            fonts_before = [(mixed.Runs(i, 1).Font.Name, mixed.Runs(i, 1).Font.Size)
                            for i in range(1, mixed.Runs().Count + 1)]
            report = fit_powerpoint_text(deck, targets)
            fonts_after = [(mixed.Runs(i, 1).Font.Name, mixed.Runs(i, 1).Font.Size)
                           for i in range(1, mixed.Runs().Count + 1)]
            self.assertEqual([f[0] for f in fonts_after], [f[0] for f in fonts_before])
            ratios = [after[1] / before[1] for before, after in zip(fonts_before, fonts_after)]
            self.assertLess(max(ratios) - min(ratios), .002)
            self.assertEqual(report[2]['lines'], 1)
            self.assertEqual(report[3]['lines'], 1)
            self.assertEqual(report[4]['lines'], 2)
            for item in report:
                self.assertLessEqual(item['width'], item['available_width'] + .2)
                self.assertLessEqual(item['height'], item['available_height'] + .2)
        finally:
            deck.Close()

    def test_original_background_pixels_survive_worker_export(self):
        from PIL import Image, ImageChops
        from pypdf import PdfReader
        from diploma_core import pptx_to_pdf
        image = Image.new('RGBA', (1200, 800), 'white')
        # Fine contrasting lines expose JPEG damage to diagrams and small text.
        for x in range(0, image.width, 11):
            for y in range(image.height):
                image.putpixel((x, y), (20, 90, 180, 255))
        source = self.root / 'background.png'
        image.save(source)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(source), 0, 0, prs.slide_width, prs.slide_height)
        text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        text.text = 'Vector text stays selectable'
        path = self.root / 'background.pptx'
        prs.save(path)
        pdf = self.root / 'background.pdf'
        pptx_to_pdf(str(path), str(pdf), threading.Event())
        page = PdfReader(pdf).pages[0]
        self.assertIn('Vector text stays selectable', page.extract_text())
        background = page.images[0].image.convert('RGBA')
        self.assertEqual(background.size, image.size)
        self.assertIsNone(ImageChops.difference(background, image).getbbox(alpha_only=False))


if __name__ == '__main__':
    unittest.main()
