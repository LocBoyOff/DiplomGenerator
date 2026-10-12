import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from datetime import datetime

from diploma_core import (Cancelled, DiplomaError, Settings, auto_mapping, fill_presentation,
                          generate, load_settings, prepare_plan, publish_pdf, read_excel,
                          read_template, safe_component, save_settings, replace_paragraph)


class GeneratorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='diploma_test_')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / 'Кириллица и пробелы'
        self.root.mkdir()
        self.excel = self.root / 'Участники.XLSX'
        self.template = self.root / 'Диплом.PPTX'
        self.settings = Settings(str(self.excel), str(self.template), str(self.root / 'Результат'))

    def workbook(self, headers, rows):
        wb = Workbook()
        wb.active.append(headers)
        for row in rows:
            wb.active.append(row)
        wb.save(self.excel)
        wb.close()

    def presentation(self, text, second=False):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        box.text = text
        if second:
            prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(self.template)
        return prs

    def row_fixture(self):
        self.settings.mode = 'row'
        self.workbook(['ФИО', 'Город', 'Дата', 'Баллы'], [['Анна', 'Тюмень', datetime(2026, 10, 12), 0]])
        self.presentation('{NAME} · {CITY} · {DATE} · {SCORE}')
        self.settings.column_mapping = {'SCORE': 'Баллы'}

    def test_row_mapping_dates_zero(self):
        self.row_fixture()
        values = prepare_plan(self.settings).records[0].values
        self.assertEqual(values['NAME'], 'Анна')
        self.assertEqual(values['CITY'], 'Тюмень')
        self.assertEqual(values['DATE'], '12.10.2026')
        self.assertEqual(values['SCORE'], '0')

    def test_team_all_people_first_trainer_and_mapping(self):
        self.workbook(['{FullN1}', 'fulln2', 'FullTr1', 'FullTr2', 'Команда', 'Город'],
                      [['Анна', 'Борис', None, 'Тренер', 'Орбита', 'Пермь']])
        self.presentation('{NAME} {TRAINER} {TEAM} {CITY}')
        plan = prepare_plan(self.settings)
        self.assertEqual(len(plan.records), 2)
        self.assertEqual([r.values['NAME'] for r in plan.records], ['Анна', 'Борис'])
        self.assertEqual(plan.records[1].values['TRAINER'], 'Тренер')
        self.assertEqual(plan.records[1].values['TEAM'], 'Орбита')

    def test_stop_on_missing(self):
        self.row_fixture()
        self.settings.column_mapping['CITY'] = ''
        with self.assertRaisesRegex(DiplomaError, 'CITY'):
            prepare_plan(self.settings)

    def test_skip_missing_record(self):
        self.settings.mode = 'row'
        self.workbook(['ФИО', 'Город'], [['Анна', None], ['Борис', 'Пермь']])
        self.presentation('{NAME} {CITY}')
        self.settings.error_handling = 'skip'
        plan = prepare_plan(self.settings)
        self.assertEqual(plan.skipped, 1)
        self.assertEqual(plan.records[0].values['NAME'], 'Борис')

    def test_default_and_unfilled_default(self):
        self.row_fixture()
        self.settings.column_mapping['CITY'] = ''
        self.settings.error_handling = 'default'
        with self.assertRaises(DiplomaError):
            prepare_plan(self.settings)
        self.settings.default_values['CITY'] = 'Не указан'
        self.assertEqual(prepare_plan(self.settings).records[0].values['CITY'], 'Не указан')

    def test_missing_team_columns(self):
        self.row_fixture()
        self.settings.mode = 'team'
        with self.assertRaisesRegex(DiplomaError, 'FullN1'):
            prepare_plan(self.settings)

    def test_duplicate_headers(self):
        self.workbook(['ФИО', '{ФИО}'], [['A', 'B']])
        with self.assertRaisesRegex(DiplomaError, 'Повторяются'):
            read_excel(self.excel)

    def test_formula_without_cache(self):
        self.workbook(['ФИО'], [['=CONCAT("A","B")']])
        with self.assertRaisesRegex(DiplomaError, 'A2'):
            read_excel(self.excel)

    def test_excel_error(self):
        self.workbook(['ФИО'], [['#DIV/0!']])
        with self.assertRaisesRegex(DiplomaError, 'ошибка Excel'):
            read_excel(self.excel)

    def test_missing_input_and_damaged_template(self):
        with self.assertRaisesRegex(DiplomaError, 'не найден'):
            read_excel(self.excel)
        self.template.write_bytes(b'broken')
        with self.assertRaisesRegex(DiplomaError, 'прочитать шаблон'):
            read_template(self.template)

    def test_stale_mapping(self):
        self.row_fixture()
        self.settings.column_mapping['NAME'] = 'Отсутствующий'
        with self.assertRaisesRegex(DiplomaError, 'отсутствующим'):
            prepare_plan(self.settings)

    def test_no_records_after_skip(self):
        self.settings.mode = 'row'
        self.workbook(['ФИО'], [['Анна']])
        self.presentation('{NAME} {CITY}')
        self.settings.error_handling = 'skip'
        with self.assertRaisesRegex(DiplomaError, 'Нет дипломов'):
            prepare_plan(self.settings)

    def test_missing_sort_field(self):
        self.row_fixture()
        self.settings.enable_sorting = True
        self.settings.sort_column = 'Absent'
        with self.assertRaisesRegex(DiplomaError, 'отсутствует'):
            prepare_plan(self.settings)

    def test_text_date_and_custom_name_mapping(self):
        self.settings.mode = 'row'
        self.workbook(['Человек', 'Дата'], [['Анна', '2026-10-12']])
        self.presentation('{NAME} {DATE}')
        self.settings.column_mapping['NAME'] = 'Человек'
        values = prepare_plan(self.settings).records[0].values
        self.assertEqual(values['DATE'], '12.10.2026')
        self.assertEqual(values['NAME'], 'Анна')

    def test_no_placeholders_or_no_slides(self):
        self.presentation('Без полей')
        with self.assertRaisesRegex(DiplomaError, 'нет полей'):
            read_template(self.template)
        Presentation().save(self.template)
        with self.assertRaisesRegex(DiplomaError, 'нет слайдов'):
            read_template(self.template)

    def test_long_destination_path(self):
        source = self.root / 'source.pdf'
        source.write_bytes(b'%PDF-1.4 test')
        long_dir = self.root / ('x' * 80) / ('y' * 80)
        with self.assertRaisesRegex(DiplomaError, 'длинный путь'):
            publish_pdf(source, long_dir, 'Имя' * 30)

    def test_copy_failure_removes_partial_pdf(self):
        source = self.root / 'source.pdf'
        source.write_bytes(b'%PDF-1.4 test')
        with patch('diploma_core.shutil.copyfileobj', side_effect=OSError('disk full')):
            with self.assertRaisesRegex(OSError, 'disk full'):
                publish_pdf(source, self.root, 'Анна')
        self.assertFalse((self.root / 'Анна.pdf').exists())

    def test_output_is_file(self):
        self.row_fixture()
        Path(self.settings.output_dir).write_bytes(b'occupied')
        with self.assertRaisesRegex(DiplomaError, 'Нельзя записать'):
            generate(self.settings, lambda x: None, lambda *x: None, threading.Event())

    def test_targeted_auto_mapping(self):
        self.assertEqual(auto_mapping(['NAME', 'CITY', 'UNKNOWN'], ['ФИО', 'Город'], 'row'),
                         {'NAME': 'ФИО', 'CITY': 'Город'})

    def test_split_repeated_tokens_preserve_style(self):
        prs = self.presentation('')
        p = prs.slides[0].shapes[0].text_frame.paragraphs[0]
        for text, bold in [('До {NA', True), ('ME} / {NAME}', False), (' после', True)]:
            run = p.add_run()
            run.text = text
            run.font.bold = bold
            run.font.size = Pt(24)
        replace_paragraph(p, {'NAME': 'Анна'})
        self.assertEqual(p.text, 'До Анна / Анна после')
        self.assertTrue(p.runs[0].font.bold)
        self.assertFalse(p.runs[1].font.bold)
        self.assertEqual(p.runs[2].text, ' после')

    def test_replacement_is_not_recursive_and_keeps_zero(self):
        prs = self.presentation('{NAME} {NAME} {COUNT}')
        p = prs.slides[0].shapes[0].text_frame.paragraphs[0]
        replace_paragraph(p, {'NAME': '{COUNT}', 'COUNT': 0})
        self.assertEqual(p.text, '{COUNT} {COUNT} 0')

    def test_soft_break_offsets(self):
        prs = self.presentation('До\v{NAME} после')
        p = prs.slides[0].shapes[0].text_frame.paragraphs[0]
        replace_paragraph(p, {'NAME': 'Анна'})
        self.assertEqual(p.text, 'До\vАнна после')

    def test_template_font_properties_unchanged_in_both_modes(self):
        for mode in ('row', 'team'):
            with self.subTest(mode=mode):
                self.settings.mode = mode
                self.workbook(['FullN1', 'CITY'], [['Анна', 'Тюмень']])
                prs = self.presentation('{NAME} {CITY}')
                run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
                run.font.name = 'Times New Roman'
                run.font.size = Pt(30)
                run.font.bold, run.font.italic, run.font.underline = True, True, True
                run.font.color.rgb = RGBColor(20, 60, 120)
                style = run._r.rPr.xml
                prs.save(self.template)
                plan = prepare_plan(self.settings)
                fill_presentation(prs, plan.records[0].values)
                self.assertEqual(run._r.rPr.xml, style)
                self.assertEqual(run.text, 'Анна Тюмень')
                self.assertFalse(prs.slides[0].shapes[0].text_frame.word_wrap)

    def test_autofit_disabled_preserves_template_wrap_and_size(self):
        prs = self.presentation('{NAME}')
        frame = prs.slides[0].shapes[0].text_frame
        frame.word_wrap = True
        before = frame._txBody.bodyPr.xml
        targets = fill_presentation(prs, {'NAME': 'Очень длинное имя'}, auto_fit=False)
        self.assertEqual(targets, [])
        self.assertEqual(frame._txBody.bodyPr.xml, before)

    def test_group_table_and_first_slide_only(self):
        prs = self.presentation('{NAME}', second=True)
        slide = prs.slides[0]
        table = slide.shapes.add_table(1, 1, Inches(0), Inches(0), Inches(3), Inches(1)).table
        table.cell(0, 0).text = '{NAME}'
        group = slide.shapes.add_group_shape()
        box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        box.text = '{NAME}'
        prs.save(self.template)
        _, placeholders = read_template(self.template)
        self.assertEqual(placeholders, ['NAME'])
        fill_presentation(prs, {'NAME': 'Анна'})
        self.assertEqual(table.cell(0, 0).text, 'Анна')
        self.assertEqual(box.text, 'Анна')
        self.assertEqual(len(prs.slides), 1)
        prs.save(self.template)
        self.assertEqual(len(Presentation(str(self.template)).slides), 1)

    def test_safe_names(self):
        for value in ['CON', 'nul.txt', 'LPT1', 'COM²']:
            self.assertTrue(safe_component(value).startswith('_'))
        self.assertEqual(safe_component('../../'), '_.._')
        self.assertEqual(safe_component('a:b?c. '), 'a_b_c')
        self.assertLessEqual(len(safe_component('Я' * 300)), 80)

    def test_publish_does_not_overwrite(self):
        source = self.root / 'source.pdf'
        source.write_bytes(b'%PDF-1.4 test')
        first = publish_pdf(source, self.root, 'Анна')
        first.write_bytes(b'original')
        second = publish_pdf(source, self.root, 'Анна')
        self.assertEqual(second.name, 'Анна (2).pdf')
        self.assertEqual(first.read_bytes(), b'original')

    def test_config_roundtrip_corrupt_and_wrong_types(self):
        path = self.root / 'config.json'
        self.settings.default_values = {'CITY': 'Пермь'}
        save_settings(self.settings, path)
        self.assertEqual(load_settings(path), self.settings)
        path.write_text('{broken', encoding='utf-8')
        with self.assertRaises(DiplomaError):
            load_settings(path)
        path.write_text(json.dumps({'enable_sorting': 'False'}), encoding='utf-8')
        with self.assertRaises(DiplomaError):
            load_settings(path)

    def test_generate_cyrillic_collision_folder_and_cleanup(self):
        self.settings.mode = 'row'
        self.workbook(['ФИО', 'Город'], [['Анна', '../CON'], ['Анна', '../CON']])
        self.presentation('{NAME} {CITY}')
        self.settings.enable_sorting = True
        self.settings.sort_column = 'CITY'
        observed = []
        def converter(pptx, pdf, stop):
            observed.append(Path(pptx).parent)
            presentation = Presentation(pptx)
            self.assertEqual(presentation.slides[0].shapes[0].text, 'Анна ../CON')
            Path(pdf).write_bytes(b'%PDF-1.4 test')
            return True
        count = generate(self.settings, lambda x: None, lambda *x: None,
                         threading.Event(), converter=converter)
        self.assertEqual(count, 2)
        pdfs = list(Path(self.settings.output_dir).rglob('*.pdf'))
        self.assertEqual(len(pdfs), 2)
        self.assertEqual({p.name for p in pdfs}, {'Анна.pdf', 'Анна (2).pdf'})
        self.assertTrue(all(p.is_relative_to(Path(self.settings.output_dir)) for p in pdfs))
        self.assertFalse(observed[0].exists())

    def test_cancellation_and_conversion_failure_cleanup(self):
        self.row_fixture()
        stop = threading.Event()
        stop.set()
        with self.assertRaises(Cancelled):
            generate(self.settings, lambda x: None, lambda *x: None, stop)
        stop.clear()
        observed = []
        def converter(pptx, pdf, event):
            observed.append(Path(pptx).parent)
            raise RuntimeError('PowerPoint недоступен')
        with self.assertRaisesRegex(DiplomaError, 'Уже сохранено дипломов: 0'):
            generate(self.settings, lambda x: None, lambda *x: None, stop, converter=converter)
        self.assertFalse(observed[0].exists())


if __name__ == '__main__':
    unittest.main()
